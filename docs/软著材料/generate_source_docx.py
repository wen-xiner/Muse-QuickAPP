# -*- coding: utf-8 -*-
"""生成软著源程序 Word 文档：前30页 + 后30页，每页50行，页眉含软件名/版本/页码"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"D:\Personal\Muse_QuickAPP"
OUT_DIR = os.path.join(ROOT, "docs", "软著材料")
OUT = os.path.join(OUT_DIR, "Muse源程序_前30页后30页.docx")
SOFT_NAME = "心屿 Muse 青少年心理健康辅助软件"
VERSION = "V1.0"
PAGE_LINES = 50
HEAD_PAGES = 30

FILES = [
    "common/constants/emotions.js",
    "common/constants/distortions.js",
    "common/constants/crisis.js",
    "common/constants/chat-scenes.js",
    "common/constants/breathing.js",
    "common/constants/backgrounds.js",
    "common/constants/backgroundThemes.js",
    "common/constants/index.js",
    "common/utils/storage.js",
    "common/utils/account.js",
    "common/utils/date.js",
    "common/utils/network.js",
    "common/utils/performance.js",
    "common/utils/resolveBackground.js",
    "common/utils/resolveOverlay.js",
    "common/utils/resolvePriority.js",
    "common/utils/reactiveBackgroundStore.js",
    "common/stores/user.js",
    "common/stores/chat.js",
    "common/stores/diary.js",
    "common/stores/index.js",
    "services/config.js",
    "services/logger.js",
    "services/ai.js",
    "functions/deepseek-proxy.js",
    "components/ChatMessage.vue",
    "components/CrisisModal.vue",
    "components/BackgroundLayer.vue",
    "components/BackgroundPicker.vue",
    "pages/login/index.vue",
    "pages/onboarding/index.vue",
    "pages/index/index.vue",
    "pages/diary/list.vue",
    "pages/diary/edit.vue",
    "pages/diary/detail.vue",
    "pages/chat/index.vue",
    "pages/chat/conversation.vue",
    "pages/regulation/index.vue",
    "pages/user/index.vue",
    "common/styles/common.scss",
]

def collect_lines():
    lines = []
    for f in FILES:
        p = os.path.join(ROOT, f)
        with open(p, encoding="utf-8") as fh:
            content = fh.read()
        lines.append("// ====== 文件: {} ======".format(f))
        lines.extend(content.splitlines())
    return lines

def set_run_font(run, size=8):
    run.font.name = "Consolas"
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Consolas")
    rfonts.set(qn("w:hAnsi"), "Consolas")
    rfonts.set(qn("w:eastAsia"), "宋体")

def add_page_number(par):
    run = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    set_run_font(run, 9)

def build():
    lines = collect_lines()
    total = len(lines)
    head = lines[:HEAD_PAGES * PAGE_LINES]
    tail = lines[-HEAD_PAGES * PAGE_LINES:] if total > HEAD_PAGES * PAGE_LINES else []
    selected = head + tail
    print("总行数: {}, 前30页: {} 行, 后30页: {} 行, 合计: {} 行".format(
        total, len(head), len(tail), len(selected)))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.0)

    hdr = sec.header
    hp = hdr.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
    hp.text = ""
    r1 = hp.add_run("{} {}".format(SOFT_NAME, VERSION))
    set_run_font(r1, 9)
    hp.add_run("\t")
    add_page_number(hp)

    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(8)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for i, ln in enumerate(selected):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(11)
        run = p.add_run(ln if ln else " ")
        set_run_font(run, 8)
        if (i + 1) % PAGE_LINES == 0 and i != len(selected) - 1:
            run.add_break(WD_BREAK.PAGE)

    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(OUT)
    print("已生成:", OUT)

if __name__ == "__main__":
    build()
